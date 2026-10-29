"""Luca dışa aktarım DÜZENİNDE sentetik yevmiye, defteri kebir ve muavin üretir (test için).

Düzen, Luca'dan alınmış gerçek örneklerin satır yapısına göre kuruldu; içerik tamamen
uydurmadır (synthetic.py'deki aynı fişler). Hesaplar üç seviyelidir: 100 > 100.01 > 100.01.001.

hatali=True olduğunda beklenen bulgular:
- K-02: Y-2026-0920 fişi yevmiyede denk değil (madde numarasıyla raporlanır)
- K-09: 100 Kasa, 14.09.2026'da Y-2026-0912 ile eksiye düşer (kebir: nakli yekün + hareketler)
- K-05K: Y-2026-0921 yalnızca yevmiyede -> 102 ve 770 ana hesaplarında yevmiye–kebir farkı
- K-05M: aynı kayıt -> 102.02.001 hesabında yevmiye–muavin farkı
kismi=True olduğunda yevmiye yalnızca 01.07–05.09 aralığını içerir (K-00).
"""

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from .synthetic import ACCOUNTS, OPENING, PERIOD, build

MUAVIN_HESAPLAR = ("100.01.001", "102.02.001")


def leaf(code: str) -> str:
    return code + ".001"


def _name(code: str) -> str:
    return ACCOUNTS.get(code) or ACCOUNTS[code.rsplit(".", 1)[0]]


def _dt(d: date) -> datetime:
    return datetime(d.year, d.month, d.day)


def _numbered(entries):
    """Fişlere madde ve fiş numarası verir (Luca'da madde numarası yevmiyede sıralıdır)."""
    out = []
    for i, (fis, d, desc, lines) in enumerate(sorted(entries, key=lambda e: (e[1], e[0]))):
        out.append({"madde": 2001 + i, "fis": 1001 + i, "ref": fis, "tarih": d, "aciklama": desc,
                    "lines": [(leaf(c), int(b * 100), int(a * 100)) for c, b, a in lines]})
    return out


def _info_rows(ws, title, start, end):
    ws.append(["SENTETİK ÖRNEK TİCARET A.Ş. — DEMO"])
    ws.append([title])
    ws.append(["Dönem :", f"{PERIOD[0]:%d/%m/%Y}-{PERIOD[1]:%d/%m/%Y}"])
    ws.append(["Tarih Aralığı :", f"{start:%d/%m/%Y}-{end:%d/%m/%Y}"])
    ws.append([])


def write_yevmiye(path: Path, fisler, start, end):
    wb = Workbook()
    ws = wb.active
    ws.title = "Yevmiye"
    _info_rows(ws, "YEVMİYE DEFTERİ", start, end)
    for f in fisler:
        if not (start <= f["tarih"] <= end):
            continue
        ws.append([f"{f['madde']:05d}-----{f['fis']:05d}-----MAHSUP-----{f['tarih']:%d/%m/%Y}"])
        ws.append(["HESAP KODU", "HESAP ADI", "AÇIKLAMA", "DETAY", "BORÇ", "ALACAK"])
        tb = ta = 0
        for side in ("borc", "alacak"):
            by_ana = defaultdict(list)
            for code, b, a in f["lines"]:
                amt = b if side == "borc" else a
                if amt:
                    by_ana[code[:3]].append((code, amt))
            for ana, items in by_ana.items():
                total = sum(x for _, x in items)
                ws.append([ana, _name(ana), None, None, total / 100 if side == "borc" else None,
                           total / 100 if side == "alacak" else None])
                for code, amt in items:
                    ws.append([code.rsplit(".", 1)[0], _name(code.rsplit(".", 1)[0]), None, amt / 100])
                    ws.append([f"   {code}", _name(code), f["aciklama"], amt / 100])
                if side == "borc":
                    tb += total
                else:
                    ta += total
        ws.append(["FİŞ AÇIKLAMA :", f["aciklama"]])
        ws.append(["TOPLAM :", None, None, None, tb / 100, ta / 100])
        ws.append([])
    wb.save(path)


def write_kebir(path: Path, fisler, start, end):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kebir"
    _info_rows(ws, "DEFTERİ KEBİR", start, end)
    ws.append(["TARİH", "MADDE NO", "FİŞ NO", "EVRAK NO", "EVRAK TARİHİ", "HESAP KODU", "HESAP ADI", "AÇIKLAMA",
               "BORÇ", "ALACAK", "BAKİYE", "B/A"])
    anas = sorted({c[:3] for c in ACCOUNTS})
    for ana in anas:
        rows = [(f, code, b, a) for f in fisler if start <= f["tarih"] <= end for code, b, a in f["lines"] if code[:3] == ana]
        op = sum(v for c, v in OPENING.items() if c[:3] == ana) * 100
        if not rows and not op:
            continue
        ws.append([f"{ana} {_name(ana)}"])
        tb, ta = max(op, 0), max(-op, 0)
        ws.append([None] * 7 + ["NAKLİ YEKÜN", tb / 100, ta / 100, abs(op) / 100, "B" if op >= 0 else "A"])
        bal = op
        for f, code, b, a in rows:
            bal += b - a
            tb += b
            ta += a
            ws.append([_dt(f["tarih"]), f["madde"], f"{f['fis']:05d}", None, None, code, _name(code), f["aciklama"],
                       b / 100 or None, a / 100 or None, abs(bal) / 100, "B" if bal >= 0 else "A"])
        ws.append(["Toplam:", None, None, None, None, None, None, None, tb / 100, ta / 100, abs(bal) / 100,
                   "B" if bal >= 0 else "A"])
    wb.save(path)


def write_muavin(path: Path, fisler, start, end, accounts=MUAVIN_HESAPLAR):
    wb = Workbook()
    ws = wb.active
    ws.title = "Muavin"
    _info_rows(ws, "MUAVİN DEFTER", start, end)
    ws.append(["TARİH", "TİP", "FİŞ NO", "AÇIKLAMA", "BORÇ", "ALACAK", "BAKİYE", "B/A"])
    gb = ga = 0
    for acc in accounts:
        op = OPENING.get(acc.rsplit(".", 1)[0], 0) * 100
        ws.append([f"{acc} {_name(acc)} TL"])
        ws.append([None, None, None, "Nakli Yekün:", max(op, 0) / 100, max(-op, 0) / 100, 0, None])
        tb = ta = 0
        for f in fisler:
            if not (start <= f["tarih"] <= end):
                continue
            for code, b, a in f["lines"]:
                if code == acc:
                    tb += b
                    ta += a
                    # Luca çıktısındaki gibi BAKİYE kolonu 0 yazılır.
                    ws.append([_dt(f["tarih"]), "Mahsup", f"{f['fis']:05d}", f["aciklama"], b / 100 or None, a / 100 or None, 0, None])
        ws.append(["Nakli Yekün Hariç", None, None, None, tb / 100, ta / 100])
        gb += tb + max(op, 0)
        ga += ta + max(-op, 0)
    ws.append(["Genel Toplam", None, None, None, gb / 100, ga / 100])
    wb.save(path)


def generate(out_dir: Path, hatali: bool = False, kismi: bool = False) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    _mizan_entries, yev_entries = build(hatali)
    yev = _numbered(yev_entries)
    # Kebir ve muavin, mizanla aynı (yalnızca yevmiyede olan Y-2026-0921 hariç) kayıtları içerir.
    # Y-2026-0920 kebirde de aynı hatalı tutarla yer alır; fark yalnızca yevmiyenin denkliğinde görülür.
    ledger = [f for f in yev if f["ref"] != "Y-2026-0921"]
    suffix = "_hatali" if hatali else ""
    paths = {
        "yevmiye": out_dir / f"DEMO_luca_yevmiye{suffix}{'_kismi' if kismi else ''}.xlsx",
        "kebir": out_dir / f"DEMO_luca_kebir{suffix}.xlsx",
        "muavin": out_dir / f"DEMO_luca_muavin{suffix}.xlsx",
    }
    yev_end = date(2026, 9, 5) if kismi else PERIOD[1]
    write_yevmiye(paths["yevmiye"], yev, PERIOD[0], yev_end)
    write_kebir(paths["kebir"], ledger, *PERIOD)
    write_muavin(paths["muavin"], ledger, *PERIOD)
    return paths
