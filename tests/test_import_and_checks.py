from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from denetci.checks.engine import run_analysis
from denetci.importer.load import DuplicateFile, ImportError_, confirm, create_draft
from denetci.models import Finding, FindingEvidence, JournalLine, RawRow, TrialBalanceLine


def _import_all(session, c, p, files):
    batches = {}
    for kind, path in files.items():
        b = create_draft(session, c, p, path.name, path.read_bytes())
        assert b.doc_type == kind
        confirm(session, b, b.doc_type, b.mapping, b.number_format)
        batches[kind] = b
    return batches


def test_detection_and_counts(session, client_period, files_hatali):
    c, p = client_period
    b = _import_all(session, c, p, files_hatali)
    mz = b["mizan"].summary["sayilar"]
    # 27 hesap satırı (ana + alt); sınıf toplamları ve genel toplam kayıt sayılmaz.
    assert mz["kayit"] == 27
    assert mz["ara_toplam"] == 6
    yv = b["yevmiye"].summary["sayilar"]
    assert yv["kayit"] == 33
    assert yv["ara_toplam"] == 3  # 2 nakli yekün + genel toplam
    assert yv["baslik"] == 3  # ilk başlık + 2 tekrar eden başlık


def test_subtotals_not_double_counted(session, client_period, files_temiz):
    c, p = client_period
    b = _import_all(session, c, p, files_temiz)
    leaves = session.scalars(select(TrialBalanceLine).where(TrialBalanceLine.batch_id == b["mizan"].id,
                                                            TrialBalanceLine.is_leaf)).all()
    total_yev = session.scalar(select(func.sum(JournalLine.borc)).where(JournalLine.batch_id == b["yevmiye"].id))
    # Ana hesaplar toplanmaz; yaprak toplamı yevmiyeyle birebir aynıdır.
    assert sum(ln.borc for ln in leaves) == total_yev
    assert all(not ln.hesap_kodu.isdigit() or len(ln.hesap_kodu) > 3 for ln in leaves)


def test_fis_no_carried_forward(session, client_period, files_temiz):
    c, p = client_period
    b = _import_all(session, c, p, files_temiz)
    lines = session.scalars(select(JournalLine).where(JournalLine.batch_id == b["yevmiye"].id,
                                                      JournalLine.fis_no == "Y-2026-0701")).all()
    assert len(lines) == 3 and all(ln.tarih == date(2026, 7, 5) for ln in lines)


def test_clean_set_has_no_findings(session, client_period, files_temiz):
    c, p = client_period
    _import_all(session, c, p, files_temiz)
    run = run_analysis(session, c, p)
    assert run.reliability == "uygun"
    assert session.scalar(select(func.count()).select_from(Finding).where(Finding.run_id == run.id)) == 0


def test_known_errors_found(session, client_period, files_hatali):
    c, p = client_period
    _import_all(session, c, p, files_hatali)
    run = run_analysis(session, c, p)
    found = {(f.rule_id, f.hesap_kodu or f.title) for f in session.scalars(select(Finding).where(Finding.run_id == run.id))}
    assert found == {
        ("K-02", "Fiş Y-2026-0920 denk değil"),
        ("K-03", "153.01"), ("K-03", "320.01"),
        ("K-04", "120"),
        ("K-05", "102.01"), ("K-05", "102.02"), ("K-05", "770.02"),
        ("K-09", "100.01"),
    }
    status = {r["id"]: r["status"] for r in run.rule_results}
    assert status["K-01"] == "gecti"
    assert run.reliability == "kritik"  # fiş denkliği bozuk


def test_negative_cash_finding_traces_to_source_row(session, client_period, files_hatali):
    c, p = client_period
    _import_all(session, c, p, files_hatali)
    run = run_analysis(session, c, p)
    f = session.scalar(select(Finding).where(Finding.run_id == run.id, Finding.rule_id == "K-09"))
    assert f.amount == 4_825_000
    assert f.severity == "Yüksek"
    assert "14.09.2026" in f.explanation and "Y-2026-0912" in f.explanation
    evs = session.scalars(select(FindingEvidence).where(FindingEvidence.finding_id == f.id)).all()
    raws = [session.get(RawRow, e.raw_row_id) for e in evs]
    # Kanıt satırı, orijinal dosyadaki aynı satırın hücreleriyle birebir eşleşir.
    ws = load_workbook(files_hatali["mizan"]).active
    mizan_raw = next(r for r in raws if r.sheet == "Mizan")
    assert ws.cell(row=mizan_raw.row_no, column=1).value == "100.01"
    yev_raw = next(r for r in raws if r.sheet == "Yevmiye")
    ws2 = load_workbook(files_hatali["yevmiye"]).active
    assert ws2.cell(row=yev_raw.row_no, column=3).value == "100.01"
    # Kasa satırı fişin ikinci satırı: fiş no hücresi boş, önceki satırdan alındığı not edilir.
    assert ws2.cell(row=yev_raw.row_no - 1, column=1).value == "Y-2026-0912"
    assert yev_raw.note == "Fiş no önceki satırdan alındı"


def test_duplicate_upload_rejected(session, client_period, files_hatali):
    c, p = client_period
    path = files_hatali["mizan"]
    create_draft(session, c, p, path.name, path.read_bytes())
    with pytest.raises(DuplicateFile):
        create_draft(session, c, p, "baska_ad.xlsx", path.read_bytes())


def test_confirm_twice_does_not_duplicate(session, client_period, files_hatali):
    c, p = client_period
    path = files_hatali["mizan"]
    b = create_draft(session, c, p, path.name, path.read_bytes())
    confirm(session, b, b.doc_type, b.mapping, b.number_format)
    n = session.scalar(select(func.count()).select_from(TrialBalanceLine))
    confirm(session, b, b.doc_type, b.mapping, b.number_format)
    assert session.scalar(select(func.count()).select_from(TrialBalanceLine)) == n


def test_new_version_keeps_old(session, client_period, files_hatali, files_temiz):
    c, p = client_period
    old = create_draft(session, c, p, "eski.xlsx", files_hatali["mizan"].read_bytes())
    confirm(session, old, "mizan", old.mapping, "tr")
    new = create_draft(session, c, p, "yeni.xlsx", files_temiz["mizan"].read_bytes())
    confirm(session, new, "mizan", new.mapping, "tr")
    assert (old.version, new.version, new.supersedes_id) == (1, 2, old.id)
    assert session.scalar(select(func.count()).select_from(TrialBalanceLine).where(TrialBalanceLine.batch_id == old.id)) > 0


def test_missing_required_columns(session, client_period, files_hatali):
    c, p = client_period
    b = create_draft(session, c, p, "m.xlsx", files_hatali["mizan"].read_bytes())
    mapping = dict(b.mapping)
    mapping.pop("borc")
    with pytest.raises(ImportError_):
        confirm(session, b, "mizan", mapping, "tr")


def test_rejects_unsupported_files(session, client_period):
    c, p = client_period
    with pytest.raises(ImportError_):
        create_draft(session, c, p, "x.pdf", b"%PDF-1.4")
    with pytest.raises(ImportError_):
        create_draft(session, c, p, "x.xlsx", b"bu bir excel degil")


def test_csv_with_invalid_amount_and_leading_zero(session, client_period, data_root):
    c, p = client_period
    csv_text = ("Hesap Kodu;Hesap Adı;Borç;Alacak;Borç Bakiye;Alacak Bakiye\n"
                "0100.01;Kasa;1.000,00;;1.000,00;\n"
                "0320.01;Satıcı;;1.000,00;;1.000,00\n"
                "0770.01;Gider;12a;;;\n"
                "TOPLAM;;1.000,00;1.000,00;;\n")
    b = create_draft(session, c, p, "m.csv", csv_text.encode("cp1254"))
    assert b.doc_type == "mizan"
    confirm(session, b, "mizan", b.mapping, "tr")
    codes = set(session.scalars(select(TrialBalanceLine.hesap_kodu).where(TrialBalanceLine.batch_id == b.id)))
    assert codes == {"0100.01", "0320.01"}
    assert b.summary["sayilar"]["hatali"] == 1 and b.summary["sayilar"]["ara_toplam"] == 1


def test_missing_document_marks_rules_not_run(session, client_period, files_temiz):
    c, p = client_period
    b = create_draft(session, c, p, "m.xlsx", files_temiz["mizan"].read_bytes())
    confirm(session, b, "mizan", b.mapping, "tr")
    run = run_analysis(session, c, p)
    status = {r["id"]: r["status"] for r in run.rule_results}
    assert status["K-02"] == "calistirilamadi" and status["K-05"] == "calistirilamadi"
    assert run.reliability == "kismen"  # eksik veriyle analiz "tamamlandı" gösterilmez
